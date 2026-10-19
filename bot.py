import os
import logging
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from telegram import (
    Update, ReplyKeyboardMarkup, ReplyKeyboardRemove, KeyboardButton
)
from telegram.ext import (
    Application, CommandHandler, MessageHandler, ConversationHandler,
    ContextTypes, filters
)

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_ID_RAW = os.getenv("ADMIN_ID", "").strip()
ALLOWED_SELLER_IDS_RAW = os.getenv("ALLOWED_SELLER_IDS", "").strip()
EXCEL_FILE = Path(os.getenv("EXCEL_FILE", "orders.xlsx"))

if not BOT_TOKEN:
    raise RuntimeError("Не задан BOT_TOKEN. Заполните файл .env.")
if not ADMIN_ID_RAW.isdigit():
    raise RuntimeError("Не задан ADMIN_ID. Укажите числовой Telegram ID администратора.")

ADMIN_ID = int(ADMIN_ID_RAW)
ALLOWED_SELLER_IDS = {
    int(value.strip())
    for value in ALLOWED_SELLER_IDS_RAW.split(",")
    if value.strip().isdigit()
}

logging.basicConfig(
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    level=logging.INFO,
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)

(
    DATE_STATE, TT_NUMBER_STATE, ADDRESS_STATE, SIM_SERIAL_STATE,
    SIM_PHONE_STATE, AMOUNT_STATE, EMPLOYEE_STATE, CONFIRM_STATE
) = range(8)

HEADERS = [
    "ID заказа", "Дата заказа", "Номер торговой точки",
    "Адрес торговой точки", "Серийный номер SIM-карты",
    "Номер SIM-карты", "Сумма заказа", "Сотрудник, который продал",
    "Кто внёс заказ (Telegram ID)", "Дата и время внесения"
]

def is_admin(user_id: int) -> bool:
    return user_id == ADMIN_ID

def is_allowed_seller(user_id: int) -> bool:
    # Если список продавцов пустой, любой пользователь, открывший личный чат с ботом,
    # может вносить заказы. Для строгого доступа заполните ALLOWED_SELLER_IDS в .env.
    return user_id == ADMIN_ID or not ALLOWED_SELLER_IDS or user_id in ALLOWED_SELLER_IDS

def ensure_workbook():
    if EXCEL_FILE.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws.append(HEADERS)
    style_sheet(ws)
    wb.save(EXCEL_FILE)

def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    widths = [12, 15, 20, 32, 28, 22, 16, 28, 25, 24]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

def append_order(order: dict, telegram_id: int) -> int:
    ensure_workbook()
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Заказы"] if "Заказы" in wb.sheetnames else wb.active
    order_id = max(1, ws.max_row)  # Header is row 1, so first order ID is 1.
    created_at = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
    ws.append([
        order_id,
        order["date"],
        order["tt_number"],
        order["address"],
        order["sim_serial"],
        order["sim_phone"],
        float(order["amount"]),
        order["employee"],
        telegram_id,
        created_at,
    ])
    style_sheet(ws)
    wb.save(EXCEL_FILE)
    return order_id

def main_keyboard(user_id: int):
    rows = [[KeyboardButton("➕ Добавить заказ")]]
    if is_admin(user_id):
        rows.append([KeyboardButton("📥 Выгрузить Excel"), KeyboardButton("📋 Последние заказы")])
    return ReplyKeyboardMarkup(rows, resize_keyboard=True)

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not update.effective_chat or update.effective_chat.type != "private":
        await update.message.reply_text("Пожалуйста, используйте бота в личном чате.")
        return
    if not is_allowed_seller(user.id):
        await update.message.reply_text("У вас нет доступа к внесению заказов. Обратитесь к администратору.")
        return
    text = (
        "Привет! Здесь можно зарегистрировать продажу.\n\n"
        "Нажмите «➕ Добавить заказ». Продавцы могут добавлять записи, "
        "а просматривать и выгружать все заказы может только администратор."
    )
    if is_admin(user.id):
        text += "\n\nКоманды администратора: /export — Excel, /orders — последние записи."
    await update.message.reply_text(text, reply_markup=main_keyboard(user.id))

async def new_order(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not is_allowed_seller(user.id):
        await update.message.reply_text("У вас нет доступа к внесению заказов.")
        return ConversationHandler.END
    context.user_data["order"] = {}
    context.user_data["order"]["date"] = date.today().strftime("%d.%m.%Y")
    keyboard = [["✅ Сегодня", "✏️ Ввести другую дату"], ["❌ Отмена"]]
    await update.message.reply_text(
        f"Дата заказа по умолчанию: {context.user_data['order']['date']}.\n"
        "Оставить сегодняшнюю дату или ввести другую?",
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True, one_time_keyboard=True),
    )
    return DATE_STATE

async def choose_date(update: Update, context: ContextTypes.DEFAULT_TYPE):
    answer = update.message.text.strip()
    if answer == "❌ Отмена":
        return await cancel(update, context)
    if answer == "✅ Сегодня":
        pass
    elif answer == "✏️ Ввести другую дату":
        await update.message.reply_text(
            "Введите дату в формате ДД.ММ.ГГГГ, например 19.10.2026.",
            reply_markup=ReplyKeyboardRemove(),
        )
        context.user_data["waiting_custom_date"] = True
        return DATE_STATE
    else:
        try:
            datetime.strptime(answer, "%d.%m.%Y")
            context.user_data["order"]["date"] = answer
            context.user_data.pop("waiting_custom_date", None)
        except ValueError:
            await update.message.reply_text("Не понял дату. Используйте ДД.ММ.ГГГГ или кнопки ниже.")
            return DATE_STATE
    if context.user_data.pop("waiting_custom_date", False):
        pass
    await update.message.reply_text("Введите номер торговой точки:", reply_markup=ReplyKeyboardRemove())
    return TT_NUMBER_STATE

async def tt_number(update: Update, context: ContextTypes.DEFAULT_TYPE):
    value = update.message.text.strip()
    if value == "❌ Отмена":
        return await cancel(update, context)
    if not value:
        await update.message.reply_text("Номер торговой точки не должен быть пустым. Введите номер:")
        return TT_NUMBER_STATE
    context.user_data["order"]["tt_number"] = value
    await update.message.reply_text("Введите адрес торговой точки:")
    return ADDRESS_STATE

async def address(update: Update, context: ContextTypes.DEFAULT_TYPE):
    value = update.message.text.strip()
    if not value:
        await update.message.reply_text("Адрес не должен быть пустым. Введите адрес:")
        return ADDRESS_STATE
    context.user_data["order"]["address"] = value
    await update.message.reply_text("Введите серийный номер SIM-карты:")
    return SIM_SERIAL_STATE

async def sim_serial(update: Update, context: ContextTypes.DEFAULT_TYPE):
    value = update.message.text.strip()
    if not value:
        await update.message.reply_text("Серийный номер не должен быть пустым. Введите его:")
        return SIM_SERIAL_STATE
    context.user_data["order"]["sim_serial"] = value
    await update.message.reply_text("Введите номер SIM-карты (номер телефона):")
    return SIM_PHONE_STATE

async def sim_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    value = update.message.text.strip()
    if not value:
        await update.message.reply_text("Номер SIM-карты не должен быть пустым. Введите его:")
        return SIM_PHONE_STATE
    context.user_data["order"]["sim_phone"] = value
    await update.message.reply_text("Введите сумму заказа, например 1500 или 1500,50:")
    return AMOUNT_STATE

async def amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    raw = update.message.text.strip().replace("₽", "").replace(" ", "").replace(",", ".")
    try:
        value = Decimal(raw)
        if value <= 0 or not value.is_finite():
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        await update.message.reply_text("Введите положительную сумму числом, например 1500 или 1500,50.")
        return AMOUNT_STATE
    context.user_data["order"]["amount"] = value
    await update.message.reply_text("Введите имя и фамилию сотрудника, который продал:")
    return EMPLOYEE_STATE

async def employee(update: Update, context: ContextTypes.DEFAULT_TYPE):
    value = update.message.text.strip()
    if not value:
        await update.message.reply_text("Укажите сотрудника, который продал:")
        return EMPLOYEE_STATE
    context.user_data["order"]["employee"] = value
    order = context.user_data["order"]
    summary = (
        "Проверьте заказ перед сохранением:\n\n"
        f"Дата: {order['date']}\n"
        f"Номер торговой точки: {order['tt_number']}\n"
        f"Адрес: {order['address']}\n"
        f"Серийный номер SIM-карты: {order['sim_serial']}\n"
        f"Номер SIM-карты: {order['sim_phone']}\n"
        f"Сумма: {order['amount']} ₽\n"
        f"Сотрудник: {order['employee']}\n\n"
        "Сохранить заказ?"
    )
    keyboard = [["✅ Сохранить", "✏️ Заполнить заново"], ["❌ Отмена"]]
    await update.message.reply_text(
        summary,
        reply_markup=ReplyKeyboardMarkup(keyboard, resize_keyboard=True, one_time_keyboard=True),
    )
    return CONFIRM_STATE

async def confirm(update: Update, context: ContextTypes.DEFAULT_TYPE):
    answer = update.message.text.strip()
    if answer == "❌ Отмена":
        return await cancel(update, context)
    if answer == "✏️ Заполнить заново":
        return await new_order(update, context)
    if answer != "✅ Сохранить":
        await update.message.reply_text("Выберите кнопку: сохранить, заполнить заново или отменить.")
        return CONFIRM_STATE
    order = context.user_data.get("order")
    if not order:
        await update.message.reply_text("Черновик не найден. Нажмите «Добавить заказ».", reply_markup=main_keyboard(update.effective_user.id))
        return ConversationHandler.END
    try:
        order_id = append_order(order, update.effective_user.id)
    except PermissionError:
        logger.exception("Не удалось записать Excel-файл: файл может быть открыт в Excel.")
        await update.message.reply_text(
            "Не удалось записать файл. Возможно, orders.xlsx открыт на компьютере. "
            "Закройте файл и попробуйте ещё раз."
        )
        return CONFIRM_STATE
    except Exception:
        logger.exception("Ошибка сохранения заказа")
        await update.message.reply_text("Произошла ошибка при сохранении. Сообщите администратору.")
        return CONFIRM_STATE

    # Send a notification to the administrator without exposing the order to other sellers.
    if update.effective_user.id != ADMIN_ID:
        try:
            await context.bot.send_message(
                chat_id=ADMIN_ID,
                text=(
                    f"🆕 Новый заказ №{order_id}\n"
                    f"Дата: {order['date']}\n"
                    f"ТТ: {order['tt_number']} — {order['address']}\n"
                    f"Сумма: {order['amount']} ₽\n"
                    f"Сотрудник: {order['employee']}"
                ),
            )
        except Exception:
            logger.exception("Не удалось отправить уведомление администратору")
    context.user_data.pop("order", None)
    await update.message.reply_text(
        f"✅ Заказ №{order_id} сохранён.",
        reply_markup=main_keyboard(update.effective_user.id),
    )
    return ConversationHandler.END

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("order", None)
    context.user_data.pop("waiting_custom_date", None)
    await update.message.reply_text(
        "Внесение заказа отменено.",
        reply_markup=main_keyboard(update.effective_user.id),
    )
    return ConversationHandler.END

async def export_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return
    ensure_workbook()
    await update.message.reply_document(
        document=EXCEL_FILE.open("rb"),
        filename=EXCEL_FILE.name,
        caption="Все сохранённые заказы."
    )

async def recent_orders(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return
    ensure_workbook()
    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb["Заказы"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    if not rows:
        await update.message.reply_text("Пока заказов нет.")
        return
    lines = ["Последние заказы:"]
    for row in rows[-10:][::-1]:
        lines.append(
            f"№{row[0]} | {row[1]} | ТТ {row[2]} | {row[6]} ₽ | {row[7]}"
        )
    await update.message.reply_text("\n".join(lines))

async def text_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    text_value = update.message.text.strip()
    if not is_allowed_seller(user_id):
        await update.message.reply_text("У вас нет доступа к этому боту.")
        return
    if text_value == "➕ Добавить заказ":
        await new_order(update, context)
    elif text_value == "📥 Выгрузить Excel":
        await export_excel(update, context)
    elif text_value == "📋 Последние заказы":
        await recent_orders(update, context)
    else:
        await update.message.reply_text("Используйте кнопки меню или команду /start.")

def main():
    ensure_workbook()
    app = Application.builder().token(BOT_TOKEN).build()

    conversation = ConversationHandler(
        entry_points=[
            CommandHandler("new", new_order),
            MessageHandler(filters.Regex(r"^➕ Добавить заказ$"), new_order),
        ],
        states={
            DATE_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_date)],
            TT_NUMBER_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, tt_number)],
            ADDRESS_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, address)],
            SIM_SERIAL_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, sim_serial)],
            SIM_PHONE_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, sim_phone)],
            AMOUNT_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, amount)],
            EMPLOYEE_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, employee)],
            CONFIRM_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, confirm)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("export", export_excel))
    app.add_handler(CommandHandler("orders", recent_orders))
    app.add_handler(conversation)
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_router))
    logger.info("Бот запущен.")
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == "__main__":
    main()
